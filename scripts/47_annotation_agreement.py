"""Phase 3 follow-up — score the human-coded 150-window annotation sample.

The coder filled the yellow columns of `annotation/emotion_windows_150.xlsx`
(sheet "CODE HERE") or exported `emotion_windows_150_coded.csv` from the HTML
form. This script joins those labels to the model labels in
`outputs/tables/p3_22_emotion_annotation_sample150.csv` and reports:

- what the coder found (category mix by corpus x domain)
- human-vs-model label agreement (accuracy, Cohen's kappa) on three subsets
- model-confidence calibration (mean model_p_max, correct vs incorrect)
- which *non-polysemous* anchors still produced wrong-sense matches

Run:

  python scripts/47_annotation_agreement.py
  python scripts/47_annotation_agreement.py --coded annotation/emotion_windows_150_coded.csv
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from analysis_config import load_analysis_config  # noqa: E402
from analysis_io import banner, item, out_path, save_summary, save_table  # noqa: E402
from plotting import apply_style, fig_heatmap, fig_stacked_bar, save_fig  # noqa: E402

SAMPLE_CSV = ROOT / "outputs" / "tables" / "p3_22_emotion_annotation_sample150.csv"
DEFAULT_XLSX = ROOT / "annotation" / "emotion_windows_150.xlsx"
DEFAULT_CSV = ROOT / "annotation" / "emotion_windows_150_coded.csv"

CATEGORIES = ["aspect_owned", "surrounding_context", "ambiguous", "wrong_anchor_match"]
EMOTIONS = ["anger", "disgust", "fear", "joy", "neutral", "sadness", "surprise"]

CATEGORY_COLORS = {
    "aspect_owned": "#59A14F",
    "surrounding_context": "#9DB2C8",
    "ambiguous": "#EDC948",
    "wrong_anchor_match": "#D4575B",
}


# ---------------------------------------------------------------------------
# loading
# ---------------------------------------------------------------------------


def _norm(value) -> str:
    if value is None or (isinstance(value, float) and np.isnan(value)):
        return ""
    return str(value).strip().lower()


def _load_coded_xlsx(path: Path) -> pd.DataFrame:
    import openpyxl

    wb = openpyxl.load_workbook(path)
    sheet = next((n for n in wb.sheetnames if n.upper().startswith("CODE")), None)
    if sheet is None:
        raise SystemExit(f"{path}: no 'CODE HERE' sheet; sheets are {wb.sheetnames}")
    ws = wb[sheet]
    header = [_norm(c.value) for c in ws[1]]
    try:
        col = {name: header.index(name) + 1 for name in ("id", "category", "emotion")}
    except ValueError as exc:
        raise SystemExit(f"{path}: 'CODE HERE' header row missing a column: {exc}")
    conf_col = next((i + 1 for i, h in enumerate(header) if h.startswith("confidence")), None)
    notes_col = next((i + 1 for i, h in enumerate(header) if h.startswith("notes")), None)

    rows = []
    for r in range(2, ws.max_row + 1):
        rid = ws.cell(r, col["id"]).value
        if rid is None or not str(rid).strip():
            continue
        rows.append(
            {
                "annotation_id": str(rid).strip(),
                "manual_category": _norm(ws.cell(r, col["category"]).value),
                "manual_emotion": _norm(ws.cell(r, col["emotion"]).value),
                "confidence_1_to_3": ws.cell(r, conf_col).value if conf_col else None,
                "coder_notes": ws.cell(r, notes_col).value if notes_col else None,
            }
        )
    return pd.DataFrame(rows)


def _load_coded_csv(path: Path) -> pd.DataFrame:
    df = pd.read_csv(path)
    required = {"annotation_id", "manual_category", "manual_emotion"}
    missing = required - set(df.columns)
    if missing:
        raise SystemExit(f"{path}: missing column(s) {sorted(missing)}")
    out = df[["annotation_id", "manual_category", "manual_emotion"]].copy()
    out["confidence_1_to_3"] = df.get("confidence_1_to_3")
    out["coder_notes"] = df.get("coder_notes")
    out["manual_category"] = out["manual_category"].map(_norm)
    out["manual_emotion"] = out["manual_emotion"].map(_norm)
    return out


def _load_coded(path: Path) -> pd.DataFrame:
    if path.suffix.lower() == ".csv":
        return _load_coded_csv(path)
    return _load_coded_xlsx(path)


# ---------------------------------------------------------------------------
# metrics
# ---------------------------------------------------------------------------


def _agreement(subset: pd.DataFrame, name: str) -> dict:
    from sklearn.metrics import accuracy_score, cohen_kappa_score

    y_true = subset["manual_emotion"]
    y_pred = subset["model_emotion"]
    return {
        "subset": name,
        "n": int(len(subset)),
        "accuracy": round(float(accuracy_score(y_true, y_pred)), 4),
        "kappa": round(float(cohen_kappa_score(y_true, y_pred, labels=EMOTIONS)), 4),
        "kappa_quadratic": round(
            float(cohen_kappa_score(y_true, y_pred, labels=EMOTIONS, weights="quadratic")), 4
        ),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Score human-coded annotation sample")
    parser.add_argument(
        "--coded",
        type=Path,
        default=None,
        help="coded .xlsx workbook or exported CSV (default: annotation/emotion_windows_150.xlsx)",
    )
    args = parser.parse_args()

    coded_path = args.coded or (DEFAULT_XLSX if DEFAULT_XLSX.exists() else DEFAULT_CSV)
    if not coded_path.exists():
        raise SystemExit(f"no coded annotation file at {coded_path}")
    if not SAMPLE_CSV.exists():
        raise SystemExit(f"missing {SAMPLE_CSV}; run scripts/22_annotation_sample.py first")

    human = _load_coded(coded_path)
    sample = pd.read_csv(SAMPLE_CSV)
    # the template ships empty manual_* columns; the coded file is authoritative
    sample = sample.drop(
        columns=[c for c in ("manual_category", "manual_emotion",
                             "confidence_1_to_3", "coder_notes") if c in sample.columns]
    )
    merged = sample.merge(human, on="annotation_id", how="left", validate="one_to_one")
    merged["manual_category"] = merged["manual_category"].fillna("")
    merged["manual_emotion"] = merged["manual_emotion"].fillna("")
    merged["emotion_match"] = merged["manual_emotion"] == merged["model_emotion"]

    uncoded = merged["manual_category"] == ""
    if uncoded.any():
        item(f"warning: {int(uncoded.sum())} rows uncoded "
             f"({merged.loc[uncoded, 'annotation_id'].tolist()})")
    coded = merged[~uncoded].copy()

    for field, allowed in (("manual_category", CATEGORIES), ("manual_emotion", EMOTIONS)):
        bad = sorted(set(coded[field]) - set(allowed) - {""})
        if bad:
            item(f"warning: unrecognised {field} value(s): {bad}")



    # --- what the coder found ------------------------------------------------
    cat_long = (
        coded.groupby(["corpus", "domain", "manual_category"])
        .size()
        .reset_index(name="n")
    )
    totals = cat_long.groupby(["corpus", "domain"])["n"].transform("sum")
    cat_long["share"] = (cat_long["n"] / totals).round(4)
    save_table(cat_long, out_path("tables", "p3", 47, "annotation_categories"))

    wrong = coded[coded["manual_category"] == "wrong_anchor_match"]
    by_anchor = (
        wrong.groupby(["anchor", "domain"])
        .size()
        .reset_index(name="n_wrong")
        .sort_values("n_wrong", ascending=False)
    )
    flagged = coded.groupby("anchor")["polysemous"].first()
    by_anchor["polysemous_flag"] = by_anchor["anchor"].map(flagged).astype(bool)
    save_table(by_anchor, out_path("tables", "p3", 47, "annotation_wrong_anchors"))

    # --- label agreement ------------------------------------------------------
    subsets = [
        ("aspect_owned", coded["manual_category"] == "aspect_owned"),
        (
            "aspect_or_context",
            coded["manual_category"].isin(["aspect_owned", "surrounding_context"]),
        ),
        ("all_coded", coded["manual_emotion"] != ""),
    ]
    agreement_rows = [
        _agreement(coded[mask], name)
        for name, mask in subsets
        if int(mask.sum()) > 0
    ]
    agreement = pd.DataFrame(agreement_rows)
    save_table(agreement, out_path("tables", "p3", 47, "annotation_agreement"))

    primary = coded[coded["manual_category"].isin(["aspect_owned", "surrounding_context"])]
    confusion = pd.crosstab(
        primary["manual_emotion"], primary["model_emotion"]
    ).reindex(index=EMOTIONS, columns=EMOTIONS, fill_value=0)
    confusion.index.name = "manual_emotion"
    save_table(confusion.reset_index(), out_path("tables", "p3", 47, "annotation_confusion"))

    # --- per-domain agreement on the primary subset --------------------------
    domain_rows = []
    for (corpus, domain), grp in primary.groupby(["corpus", "domain"]):
        domain_rows.append(
            {
                "corpus": corpus,
                "domain": domain,
                "n": len(grp),
                "accuracy": round(float((grp["emotion_match"]).mean()), 4),
                "aspect_owned_share_of_stratum": round(
                    float((coded[(coded["corpus"] == corpus) & (coded["domain"] == domain)]
                           ["manual_category"] == "aspect_owned").mean()), 4
                ),
            }
        )
    save_table(pd.DataFrame(domain_rows), out_path("tables", "p3", 47, "annotation_by_stratum"))

    # --- calibration ----------------------------------------------------------
    calib = {}
    if "model_p_max" in primary.columns:
        from scipy.stats import mannwhitneyu

        ok = primary["model_p_max"][primary["emotion_match"]]
        bad = primary["model_p_max"][~primary["emotion_match"]]
        u_stat, u_p = mannwhitneyu(ok, bad, alternative="greater")
        calib = {
            "p_max_correct_mean": round(float(ok.mean()), 4),
            "p_max_incorrect_mean": round(float(bad.mean()), 4),
            "mannwhitney_u_p": float(f"{u_p:.3e}"),
        }

    conf = (
        primary.assign(conf=pd.to_numeric(primary["confidence_1_to_3"], errors="coerce"))
        .groupby("conf")["emotion_match"]
        .agg(["mean", "count"])
        .reset_index()
        .rename(columns={"mean": "accuracy", "count": "n"})
    )
    conf["accuracy"] = conf["accuracy"].round(4)
    save_table(conf, out_path("tables", "p3", 47, "annotation_by_confidence"))

    # non-neutral model rate by category: how much signal sits on unusable rows
    noise = (
        coded.assign(model_non_neutral=coded["model_emotion"] != "neutral")
        .groupby("manual_category")["model_non_neutral"]
        .agg(["mean", "count"])
        .reset_index()
        .rename(columns={"mean": "model_non_neutral_share", "count": "n"})
    )
    noise["model_non_neutral_share"] = noise["model_non_neutral_share"].round(4)

    save_table(
        merged[
            [
                "annotation_id", "corpus", "domain", "anchor", "polysemous",
                "manual_category", "manual_emotion", "confidence_1_to_3",
                "model_emotion", "model_p_max", "emotion_match", "post_id",
            ]
        ],
        out_path("tables", "p3", 47, "annotation_coding"),
    )

    # --- figures --------------------------------------------------------------
    cfg = load_analysis_config()
    style = apply_style(cfg.section("figures"))

    strata = sorted(coded.groupby(["corpus", "domain"]).groups)
    labels = [f"{c}/{d}" for c, d in strata]
    layers = {}
    for cat in CATEGORIES:
        shares = []
        for c, d in strata:
            grp = coded[(coded["corpus"] == c) & (coded["domain"] == d)]
            shares.append(float((grp["manual_category"] == cat).mean()))
        layers[cat] = shares
    fig, _ = fig_stacked_bar(
        labels, layers, style=style,
        title="Human coding of 150 emotion windows (share of stratum)",
        ylabel="share of windows",
        colors=CATEGORY_COLORS,
    )
    save_fig(fig, out_path("figures", "p3", 47, "annotation_categories", ext="png"))

    fig, _ = fig_heatmap(
        confusion.to_numpy(dtype=float), EMOTIONS, EMOTIONS, style=style,
        title="Human vs model emotion (aspect-owned + context windows)",
        cbar_label="windows", cmap="Blues", annotate=True,
    )
    save_fig(fig, out_path("figures", "p3", 47, "annotation_confusion", ext="png"))

    # --- summary --------------------------------------------------------------
    cat_share = (coded["manual_category"].value_counts(normalize=True).round(4)).to_dict()
    summary = {
        "source": str(coded_path.relative_to(ROOT)),
        "n_coded": int(len(coded)),
        "n_uncoded": int(uncoded.sum()),
        "category_share": cat_share,
        "category_counts": coded["manual_category"].value_counts().to_dict(),
        "agreement": agreement_rows,
        "calibration": calib,
        "wrong_anchor_match": {
            "n": int(len(wrong)),
            "share": round(float(len(wrong) / len(coded)), 4),
            "on_polysemous_flagged_anchors": int(wrong["polysemous"].sum()),
            "top_anchors": by_anchor.head(10).set_index("anchor")["n_wrong"].to_dict(),
        },
        "model_non_neutral_by_category": noise.set_index("manual_category")[
            "model_non_neutral_share"
        ].to_dict(),
        "aspect_owned_emotion_mix": coded[coded["manual_category"] == "aspect_owned"][
            "manual_emotion"
        ].value_counts().to_dict(),
        "privacy": "Restricted: row-level coding table contains post identifiers; stays local.",
    }
    save_summary(summary, out_path("summaries", "p3", 47, "annotation_agreement", ext="json"))

    banner("Annotation agreement")
    item(f"coded {len(coded)}/150 from {coded_path.name}")
    for row in agreement_rows:
        item(f"{row['subset']:20s} n={row['n']:3d}  acc={row['accuracy']:.3f}  "
             f"kappa={row['kappa']:.3f}")
    item(f"aspect_owned share: {cat_share.get('aspect_owned', 0):.1%}; "
         f"wrong_anchor_match: {cat_share.get('wrong_anchor_match', 0):.1%}")
    if calib:
        item(f"model p_max correct {calib['p_max_correct_mean']:.2f} vs "
             f"incorrect {calib['p_max_incorrect_mean']:.2f} "
             f"(MWU p={calib['mannwhitney_u_p']})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
